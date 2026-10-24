"""Build editable DOCX and matching PDF reports for laboratories 1 and 2."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
FONT = Path(r"C:\Windows\Fonts\arial.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\arialbd.ttf")


def setup_fonts():
    pdfmetrics.registerFont(TTFont("Arial", str(FONT)))
    pdfmetrics.registerFont(TTFont("Arial-Bold", str(FONT_BOLD)))


def code(path, start, end):
    lines = (ROOT / path).read_text(encoding="utf-8").splitlines()
    return "\n".join(lines[start - 1:end])


LAB01 = {
    "number": "1",
    "title": "Работа с портами микроконтроллера и битами",
    "report_dir": ROOT / "labs" / "lab01" / "reports",
    "source": "labs/lab01/ports/ports.ino",
    "goal": "Изучить настройку портов ввода-вывода ATmega328P, побитовые операции и программную обработку кнопок.",
    "theory": [
        "У ATmega328P направление вывода задаётся регистром DDRx: единица означает выход, ноль — вход. Регистр PORTx устанавливает логический уровень на выходе либо включает внутреннюю подтяжку входа. Регистр PINx позволяет прочитать фактическое состояние входного порта.",
        "В работе задействованы PD2–PD4 как входы с внутренней подтяжкой и PB3–PB5 как выходы. При нажатой кнопке вход имеет уровень LOW, поэтому прочитанные биты инвертируются. Логические биты 0–2 далее соответствуют трём кнопкам и трём светодиодам.",
    ],
    "wiring": [
        ["Элемент", "Вывод Arduino Uno / ATmega328P", "Режим"],
        ["Кнопка 1", "D2 / PD2", "вход, INPUT_PULLUP"],
        ["Кнопка 2", "D3 / PD3", "вход, INPUT_PULLUP"],
        ["Кнопка 3", "D4 / PD4", "вход, INPUT_PULLUP"],
        ["Красный LED", "D11 / PB3", "выход через 220 Ом"],
        ["Жёлтый LED", "D12 / PB4", "выход через 220 Ом"],
        ["Зелёный LED", "D13 / PB5", "выход через 220 Ом"],
    ],
    "steps": [
        "В setup() прямыми операциями над DDRD, PORTD и DDRB настроены три входа и три выхода. Это демонстрирует работу с регистрами, а не только вызовами pinMode().",
        "В loop() байт PIND читается целиком. Функция logicalPressedFromPind() инвертирует активный ноль и выделяет PD2–PD4 маской 0x07.",
        "Для светодиодов поочерёдно применяются два способа: masked read-modify-write регистра PORTB и bitSet()/bitClear(). Оба способа не затрагивают PB0–PB2.",
        "Функции firstSetLoop() и firstSetCtz() находят младший установленный бит разными способами. В строке POS печатаются все установленные позиции, поэтому одновременное нажатие не сводится к одному индексу.",
        "Антидребезг выполнен автоматом candidate/stable с интервалом 40 мс. Вывод Serial и обновление светодиодов происходят с периодом 200 мс, то есть с фактической целевой частотой 5 Гц.",
    ],
    "listing": code("labs/lab01/ports/ports.ino", 4, 32),
    "listing_caption": "Настройка портов и ключевые побитовые функции",
    "results": [
        ["Проверка", "Наблюдаемый результат", "Источник"],
        ["Сборка Uno", "2712 Б Flash (8%), 199 Б SRAM (9%)", "artifacts/build-logs/lab01-ports-20261024-141530.log"],
        ["Нет нажатий", "PRESS=00000000, POS=[none], оба FIRST=-1", "results/ports-serial.log"],
        ["Одна кнопка D2", "PRESS=00000001, POS=[0], LED=00000001", "results/ports-serial.log"],
        ["D3 и D4 одновременно", "PRESS=00000110, POS=[1,2], оба FIRST=1", "results/ports-serial.log"],
        ["Антидребезг", "10-мс импульс не изменил stable; устойчивый D2 принят", "results/ports-debounce-serial.log"],
        ["Частота", "MS=200, 400, 600, 800; интервал 200 мс = 5 Гц", "results/ports-frequency-serial.log"],
    ],
    "analysis": "Реальная Wokwi-симуляция подтвердила поведение для нулевой, одиночной и множественной масок. Строка POS=[1,2] непосредственно закрывает требование о выводе всех позиций установленных битов. Из четырёх соседних временных меток получен период 200 мс; следовательно, частота равна 1000/200 = 5 Гц.",
    "conclusion": "Цель достигнута: входной порт читается побитово, выходной порт изменяется двумя способами, поиск первого установленного бита реализован двумя методами, а программный антидребезг и периодический вывод подтверждены симуляцией.",
}

LAB02 = {
    "number": "2",
    "title": "Разработка конечного автомата двумя методами",
    "report_dir": ROOT / "labs" / "lab02" / "reports",
    "source": "labs/lab02/switch_fsm/switch_fsm.ino; labs/lab02/table_fsm/table_fsm.ino",
    "goal": "Разработать контроллер светофора как конечный автомат и сравнить реализации через switch-case и таблицу обработчиков.",
    "theory": [
        "Конечный автомат хранит текущее состояние и меняет его по событиям. Для светофора нормальные состояния GREEN, YELLOW и RED переходят друг в друга по истечению интервала. Неблокирующий таймер millis() позволяет обрабатывать кнопки во время ожидания.",
        "В варианте switch_fsm переход выбирается оператором switch(state). В table_fsm состояние и событие индексируют двухмерную таблицу указателей на обработчики. Общая логика аварийного и ночного режимов приостанавливает нормальный автомат и затем восстанавливает сохранённый остаток времени.",
    ],
    "wiring": [
        ["Элемент", "Вывод Arduino Uno", "Назначение"],
        ["Красный LED", "D13", "состояние RED"],
        ["Жёлтый LED", "D12", "состояние YELLOW; мигание в special mode"],
        ["Зелёный LED", "D11", "состояние GREEN"],
        ["Кнопка пешехода", "D2", "очередь продления красного"],
        ["Аварийная кнопка", "D3", "аварийный режим с приоритетом"],
        ["Ночной режим", "D4", "мигающий жёлтый; режим из набора тестов методички"],
    ],
    "steps": [
        "Нормальный цикл задан длительностями GREEN=10 с, YELLOW=3 с, RED=10 с. Пешеходный запрос, поступивший до красного, увеличивает RED на 5 с.",
        "Кнопки D2 и D4 обрабатываются с debounce 50 мс. Аварийный вход D3 имеет высший приоритет: запоминаются состояние и остаток времени, затем включается жёлтый сигнал; после снятия аварии нормальный автомат продолжает прерванную фазу.",
        "Ночной режим реализован как дополнительный режим с мигающим жёлтым светодиодом и тем же механизмом сохранения/восстановления. Он указан явно, потому что присутствует в тестах, хотя в методичке назван опциональным.",
        "Для ускоренного теста предусмотрена последовательная команда F: длительности сокращаются до 1 с, 0,3 с, 1 с и 0,5 с добавки. Нормальные значения остаются значениями запуска программы.",
        "Обе реализации имеют отдельные diagram.json, wokwi.toml и независимую сборку. Отличается только диспетчеризация нормальных событий.",
    ],
    "listing": code("labs/lab02/table_fsm/table_fsm.ino", 109, 135),
    "listing_caption": "Таблица обработчиков событий конечного автомата",
    "results": [
        ["Проверка", "switch-case", "таблица обработчиков"],
        ["Сборка Uno", "4400 Б Flash (13%), 422 Б SRAM (20%)", "4460 Б Flash (13%), 434 Б SRAM (21%)"],
        ["Полный нормальный цикл", "0: GREEN 10000; 10000: YELLOW 3000; 13000: RED 10000", "та же последовательность"],
        ["Аварийный режим", "EMERGENCY_ON; EMERGENCY_OFF_RESTORE, остаток 9963 мс", "та же последовательность"],
        ["Ночной режим", "NIGHT_ON; NIGHT_OFF_RESTORE", "та же последовательность"],
        ["Пешеход / debounce", "один PED_QUEUED; RED=1500 мс в fast mode", "один PED_QUEUED; RED=1500 мс в fast mode"],
    ],
    "analysis": "Финальная сборка от 24.10.2026 записана в artifacts/build-logs/lab02-switch_fsm-20261024-141534.log и lab02-table_fsm-20261024-141537.log. Успешные Wokwi-сценарии normal и events записаны отдельно для обеих версий. Логи ped/debounce содержат ожидаемые строки, но запускатель Wokwi завершил этот сценарий кодом timeout 42 уже после записи результата. Поэтому этот факт трактуется как наблюдение в логе, а не как успешный завершённый сценарий. Табличная версия дороже switch-case на 60 Б Flash и 12 Б SRAM; это ожидаемая цена таблицы указателей и отдельных обработчиков.",
    "conclusion": "Цель достигнута: один и тот же светофорный автомат реализован двумя способами, управляется неблокирующими таймерами и обрабатывает пешеходное, аварийное и ночное события. Полный цикл и специальные режимы подтверждены реальными Wokwi-логами; ограничение timeout 42 задокументировано без подмены результата.",
}


def doc_title(doc, lab):
    for text, size, bold in [
        ("Министерство науки и высшего образования Российской Федерации", 12, False),
        ("ФГБОУ ВО «Самарский государственный технический университет»", 12, True),
        ("", 12, False),
        ("ОТЧЁТ", 16, True),
        (f"по лабораторной работе № {lab['number']}", 14, True),
        (f"«{lab['title']}»", 14, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
    doc.add_paragraph("\n\n\n\n\n")
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    info.add_run("Дисциплина: Программирование микропроцессоров\n")
    info.add_run("Выполнил(а): ______________________________\n")
    info.add_run("Группа: ____________________________________\n")
    info.add_run("Преподаватель: _____________________________\n")
    doc.add_paragraph("\n\nСамара — 2026").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def doc_table(doc, rows):
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                if row_index == 0:
                    run.bold = True
    return table


def build_docx(lab):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)
    doc_title(doc, lab)
    doc.add_heading("Цель работы", level=1)
    doc.add_paragraph(lab["goal"])
    doc.add_heading("Теоретические сведения", level=1)
    for item in lab["theory"]:
        doc.add_paragraph(item)
    doc.add_heading("Схема подключения", level=1)
    doc.add_paragraph("Схема собрана в Wokwi; соответствие компонентов выводам приведено в таблице.")
    doc_table(doc, lab["wiring"])
    doc.add_heading("Выполнение работы", level=1)
    for index, item in enumerate(lab["steps"], 1):
        doc.add_paragraph(item, style="List Number")
    doc.add_heading("Ключевой фрагмент программы", level=1)
    p = doc.add_paragraph()
    p.add_run(lab["listing_caption"] + "\n").italic = True
    code_p = doc.add_paragraph()
    code_p.style = doc.styles["Normal"]
    for run in code_p.runs:
        run.font.name = "Consolas"
        run.font.size = Pt(8)
    code_p.add_run(lab["listing"])
    doc.add_paragraph("Полный исходный код расположен в " + lab["source"] + ".")
    doc.add_heading("Результаты экспериментов", level=1)
    doc_table(doc, lab["results"])
    doc.add_heading("Анализ результатов", level=1)
    doc.add_paragraph(lab["analysis"])
    doc.add_heading("Заключение", level=1)
    doc.add_paragraph(lab["conclusion"])
    doc.add_heading("Воспроизводимость", level=1)
    doc.add_paragraph("Сборка: .\\scripts\\build.ps1 -Lab lab" + lab["number"].zfill(2) + " -Sketch " + ("ports" if lab["number"] == "1" else "switch_fsm") + ".")
    if lab["number"] == "2":
        doc.add_paragraph("Табличный вариант: .\\scripts\\build.ps1 -Lab lab02 -Sketch table_fsm.")
    doc.add_paragraph("Интерактивная схема открывается через соответствующую папку в VS Code и команду Wokwi: Start Simulator.")
    out = lab["report_dir"] / f"lab{lab['number'].zfill(2)}-report.docx"
    doc.save(out)


def styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("TitleRu", parent=base["Title"], fontName="Arial-Bold", fontSize=16, leading=20, alignment=TA_CENTER),
        "subtitle": ParagraphStyle("SubtitleRu", parent=base["Normal"], fontName="Arial-Bold", fontSize=13, leading=17, alignment=TA_CENTER),
        "head": ParagraphStyle("HeadRu", parent=base["Heading1"], fontName="Arial-Bold", fontSize=14, leading=17, spaceBefore=10, spaceAfter=7),
        "body": ParagraphStyle("BodyRu", parent=base["BodyText"], fontName="Arial", fontSize=10.5, leading=14, alignment=TA_JUSTIFY, spaceAfter=7),
        "small": ParagraphStyle("SmallRu", parent=base["BodyText"], fontName="Arial", fontSize=8, leading=10),
        "caption": ParagraphStyle("CaptionRu", parent=base["BodyText"], fontName="Arial", fontSize=9, leading=11, alignment=TA_CENTER, spaceAfter=6),
        "code": ParagraphStyle("CodeRu", parent=base["Code"], fontName="Courier", fontSize=6.3, leading=7.5),
    }


def pdf_table(rows, style):
    converted = [[Paragraph(str(value), style["small"]) for value in row] for row in rows]
    width = 17.5 * cm
    table = Table(converted, colWidths=[width / len(rows[0])] * len(rows[0]), repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#6C7A89")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Arial", 8)
    canvas.drawCentredString(A4[0] / 2, 1.2 * cm, f"Лабораторная работа № {doc.lab_number} — {canvas.getPageNumber()}")
    canvas.restoreState()


def build_pdf(lab):
    setup_fonts()
    style = styles()
    out = lab["report_dir"] / f"lab{lab['number'].zfill(2)}-report.pdf"
    document = SimpleDocTemplate(str(out), pagesize=A4, rightMargin=1.5 * cm, leftMargin=2.5 * cm, topMargin=1.8 * cm, bottomMargin=1.8 * cm)
    document.lab_number = lab["number"]
    story = []
    for line in [
        "Министерство науки и высшего образования Российской Федерации",
        "ФГБОУ ВО «Самарский государственный технический университет»",
    ]:
        story.append(Paragraph(line, style["body"]))
    story.extend([Spacer(1, 2.8 * cm), Paragraph("ОТЧЁТ", style["title"]), Paragraph(f"по лабораторной работе № {lab['number']}", style["subtitle"]), Paragraph(f"«{lab['title']}»", style["subtitle"]), Spacer(1, 4.5 * cm)])
    story.append(Paragraph("Дисциплина: Программирование микропроцессоров<br/>Выполнил(а): ______________________________<br/>Группа: ____________________________________<br/>Преподаватель: _____________________________", style["body"]))
    story.extend([Spacer(1, 2 * cm), Paragraph("Самара — 2026", style["caption"]), PageBreak()])
    story.append(Paragraph("Цель работы", style["head"]))
    story.append(Paragraph(lab["goal"], style["body"]))
    story.append(Paragraph("Теоретические сведения", style["head"]))
    for item in lab["theory"]:
        story.append(Paragraph(item, style["body"]))
    story.append(Paragraph("Схема подключения", style["head"]))
    story.append(Paragraph("Собранная в Wokwi схема использует Arduino Uno. Соединения компонентов приведены в таблице.", style["body"]))
    story.append(pdf_table(lab["wiring"], style))
    story.append(Spacer(1, 0.3 * cm))
    story.append(Paragraph("Выполнение работы", style["head"]))
    for index, item in enumerate(lab["steps"], 1):
        story.append(Paragraph(f"{index}. {item}", style["body"]))
    story.append(Paragraph("Ключевой фрагмент программы", style["head"]))
    story.append(Paragraph(lab["listing_caption"], style["caption"]))
    story.append(Preformatted(lab["listing"], style["code"], maxLineLength=115))
    story.append(Paragraph("Полные исходные файлы находятся в репозитории: " + lab["source"], style["body"]))
    story.append(Paragraph("Результаты экспериментов", style["head"]))
    story.append(pdf_table(lab["results"], style))
    story.append(Paragraph("Анализ результатов", style["head"]))
    story.append(Paragraph(lab["analysis"], style["body"]))
    story.append(Paragraph("Заключение", style["head"]))
    story.append(Paragraph(lab["conclusion"], style["body"]))
    story.append(Paragraph("Воспроизводимость", style["head"]))
    sketch = "ports" if lab["number"] == "1" else "switch_fsm"
    story.append(Paragraph(f"Сборка: .\\scripts\\build.ps1 -Lab lab{lab['number'].zfill(2)} -Sketch {sketch}. Интерактивная схема открывается в VS Code из соответствующей папки командой Wokwi: Start Simulator.", style["body"]))
    if lab["number"] == "2":
        story.append(Paragraph("Для табличного варианта используется команда .\\scripts\\build.ps1 -Lab lab02 -Sketch table_fsm.", style["body"]))
    document.build(story, onFirstPage=footer, onLaterPages=footer)


if __name__ == "__main__":
    for lab in (LAB01, LAB02):
        lab["report_dir"].mkdir(parents=True, exist_ok=True)
        build_docx(lab)
        build_pdf(lab)
        print(f"Generated lab{lab['number'].zfill(2)} report")
